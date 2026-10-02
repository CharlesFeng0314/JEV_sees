"""Behavior that does not need a camera model or a network call."""

from __future__ import annotations

import json
import unittest

import numpy as np
from PIL import Image
from typesafe_sdk import Choice, Noul

from jev_sees.color import cv_color_evidence
from jev_sees.given_that import build_given_that
from jev_sees.memory import SceneMemory
from jev_sees.render import DashboardRenderer
from jev_sees.report import JudgmentLog, write_excel, write_json
from jev_sees.result import Result
from jev_sees.runtime import DENSE_REGION_CAPTION, VisionRuntime, _dense_regions
from jev_sees.session import Sees
from jev_sees.tracking import TrackBank, box_gap, box_iou


class PerceptionTests(unittest.TestCase):
    def test_cv_color_evidence_keeps_measurements(self) -> None:
        blue = np.zeros((4, 5, 3), dtype=np.uint8)
        blue[:, :, 2] = 255

        evidence = cv_color_evidence(blue)

        self.assertEqual(evidence["label"], "blue")
        self.assertEqual(evidence["median_rgb"], [0, 0, 255])
        self.assertEqual(evidence["pixel_count"], 20)
        self.assertEqual(evidence["chromatic_fraction"], 1.0)

    def test_clip_color_evidence_keeps_the_distribution(self) -> None:
        runtime = VisionRuntime()

        def classify(crops, labels):
            return [
                {
                    label: 0.72 if "blue" in label else 0.28 if "green" in label else 0.0
                    for label in labels
                }
            ]

        runtime.classify_crops = classify
        evidence = runtime.color_evidence([Image.fromarray(np.zeros((2, 2, 3), dtype=np.uint8))])

        self.assertEqual(evidence[0]["label"], "blue")
        self.assertEqual(evidence[0]["confidence"], 0.72)
        self.assertEqual(evidence[0]["probabilities"]["green"], 0.28)

    def test_florence_dense_regions_need_no_object_vocabulary(self) -> None:
        parsed = {
            DENSE_REGION_CAPTION: {
                "bboxes": [[-2, 3, 30, 40], [9, 9, 2, 2], [1, 1, 1, 5]],
                "labels": ["  a blue bus. ", "person", "invalid"],
            }
        }

        regions = _dense_regions(parsed, (20, 10))

        self.assertEqual(
            regions,
            [
                {
                    "label": "a blue bus",
                    "description": "a blue bus",
                    "bbox_xyxy": [0.0, 3.0, 20.0, 10.0],
                },
                {
                    "label": "person",
                    "description": "person",
                    "bbox_xyxy": [2.0, 2.0, 9.0, 9.0],
                },
            ],
        )

    def test_sees_accepts_an_autonomous_perceptor(self) -> None:
        class Perceptor:
            def detect(self, rgb):
                return [
                    {
                        "label": "a city bus",
                        "description": "a blue city bus",
                        "bbox_xyxy": [1, 1, 7, 7],
                    }
                ]

            def color_names(self, crops):
                return ["blue"]

        tracks = Sees(perceptor=Perceptor()).observe(np.zeros((8, 8, 3), dtype=np.uint8))

        self.assertEqual(tracks[0]["label"], "a city bus")
        self.assertEqual(tracks[0]["object_id"], "object_001")
        evidence = tracks[0]["attributes"]["color_evidence"]
        self.assertEqual(set(evidence), {"cv", "clip", "caption"})
        self.assertEqual(evidence["caption"]["text"], "a blue city bus")


class TrackingTests(unittest.TestCase):
    def test_same_box_keeps_id(self) -> None:
        bank = TrackBank("image")
        first = bank.update(
            [{"label": "bus", "bbox_xyxy": [10, 10, 80, 80], "centroid_uv": [45, 45]}]
        )
        second = bank.update(
            [{"label": "bus", "bbox_xyxy": [12, 11, 82, 81], "centroid_uv": [47, 46]}]
        )
        self.assertEqual(first[0]["object_id"], second[0]["object_id"])

    def test_distant_box_is_a_new_id(self) -> None:
        bank = TrackBank("image")
        first = bank.update([{"label": "bus", "bbox_xyxy": [10, 10, 40, 40]}])
        second = bank.update([{"label": "car", "bbox_xyxy": [400, 400, 500, 500]}])
        self.assertNotEqual(first[0]["object_id"], second[0]["object_id"])

    def test_nearby_non_overlapping_box_keeps_id(self) -> None:
        bank = TrackBank("image")
        first = bank.update([{"label": "person", "bbox_xyxy": [10, 10, 40, 90]}])
        second = bank.update([{"label": "person", "bbox_xyxy": [42, 10, 72, 90]}])
        self.assertEqual(first[0]["object_id"], second[0]["object_id"])

    def test_nearby_different_label_does_not_keep_id_without_overlap(self) -> None:
        bank = TrackBank("image")
        first = bank.update([{"label": "person", "bbox_xyxy": [10, 10, 40, 90]}])
        second = bank.update([{"label": "car", "bbox_xyxy": [42, 10, 72, 90]}])
        self.assertNotEqual(first[0]["object_id"], second[0]["object_id"])

    def test_camera_position_keeps_id(self) -> None:
        bank = TrackBank("camera")
        first = bank.update(
            [{"label": "cup", "position_m": [0.1, 0.2, 0.6], "bbox_xyxy": [0, 0, 10, 10]}]
        )
        second = bank.update(
            [{"label": "cup", "position_m": [0.12, 0.21, 0.61], "bbox_xyxy": [1, 1, 11, 11]}]
        )
        self.assertEqual(first[0]["object_id"], second[0]["object_id"])

    def test_overlap_and_gap(self) -> None:
        self.assertGreater(box_iou([0, 0, 10, 10], [5, 5, 15, 15]), 0)
        self.assertEqual(box_gap([0, 0, 10, 10], [12, 0, 20, 10]), 2)


class MemoryTests(unittest.TestCase):
    def test_label_vote_and_stale(self) -> None:
        memory = SceneMemory()
        memory.observe(
            [{"object_id": "object_001", "label": "bus", "confidence": 0.4, "bbox_xyxy": [0, 0, 10, 10]}]
        )
        memory.observe(
            [{"object_id": "object_001", "label": "bus", "confidence": 0.8, "bbox_xyxy": [1, 1, 11, 11]}]
        )
        memory.observe(
            [{"object_id": "object_001", "label": "truck", "confidence": 0.3, "bbox_xyxy": [1, 1, 11, 11]}]
        )
        self.assertEqual(memory.known_objects["object_001"]["label"], "bus")
        for _ in range(3):
            memory.observe([])
        record = memory.known_objects["object_001"]
        self.assertFalse(record["currently_visible"])
        self.assertTrue(record["stale"])
        self.assertLessEqual(len(record["pose_history"]), 20)

    def test_video_time_reuses_pose_history(self) -> None:
        memory = SceneMemory()
        for frame in (10, 12, 14):
            memory.observe(
                [
                    {
                        "object_id": "object_001",
                        "label": "person",
                        "bbox_xyxy": [frame, 1, frame + 4, 9],
                        "centroid_uv": [frame + 2, 5],
                    }
                ],
                frame_index=frame,
                video_time_s=frame / 4,
            )
        record = memory.known_objects["object_001"]
        self.assertEqual(len(record["pose_history"]), 3)
        projected = memory.prompt_objects()[0]["recent_poses"]
        self.assertEqual([item["frame_index"] for item in projected], [12, 14])
        self.assertEqual(projected[-1]["video_time_s"], 3.5)


class GivenThatTests(unittest.TestCase):
    def test_facts_keep_pixels_and_drop_scores(self) -> None:
        memory = SceneMemory()
        memory.observe(
            [
                {
                    "object_id": "object_001",
                    "label": "bus",
                    "confidence": 0.91,
                    "bbox_xyxy": [10, 20, 100, 80],
                    "centroid_uv": [55, 50],
                    "attributes": {
                        "color_evidence": {
                            "cv": {
                                "label": "yellow",
                                "median_rgb": [220, 190, 30],
                                "chromatic_fraction": 0.8,
                            },
                            "clip": {
                                "label": "yellow",
                                "confidence": 0.75,
                                "probabilities": {"yellow": 0.75, "orange": 0.25},
                            },
                            "caption": {"text": "a yellow bus"},
                        }
                    },
                    "semantic_scores": {"bus": 0.9, "car": 0.1},
                    "description": "yellow bus",
                },
                {
                    "object_id": "object_002",
                    "label": "person",
                    "confidence": 0.5,
                    "bbox_xyxy": [90, 20, 130, 90],
                    "centroid_uv": [110, 55],
                    "attributes": {"color": "blue"},
                },
            ]
        )
        state = build_given_that(
            "What color is the bus?", memory, modality="rgb", image_size=(640, 480)
        )
        payload = json.dumps(state)
        self.assertNotIn("semantic_scores", payload)
        self.assertNotIn("pose_history", payload)
        visible = state["current_scene"]["visible_objects"]
        self.assertEqual(visible[0]["bbox_xyxy"], [10, 20, 100, 80])
        self.assertEqual(state["schema_version"], 2)
        color = visible[0]["attributes"]["color_evidence"]
        self.assertEqual(set(color), {"cv", "clip", "caption"})
        self.assertEqual(color["cv"]["median_rgb"], [220, 190, 30])
        self.assertEqual(color["clip"]["probabilities"]["yellow"], 0.75)
        self.assertEqual(color["caption"]["text"], "a yellow bus")
        self.assertNotIn("color", visible[0]["attributes"])
        self.assertTrue(state["relations"])
        self.assertIn("prompt_budget", state)
        self.assertEqual(state["prompt_budget"]["hard_limit_tokens"], 31000)
        self.assertLessEqual(state["prompt_budget"]["serialized_chars"], 24000)
        self.assertEqual(
            state["prompt_budget"]["serialized_chars"],
            len(json.dumps(state, ensure_ascii=False, separators=(",", ":"), default=str)),
        )

    def test_current_scene_exposes_recent_video_poses_only(self) -> None:
        memory = SceneMemory()
        for frame in (2, 4):
            memory.observe(
                [
                    {
                        "object_id": "object_001",
                        "label": "person",
                        "bbox_xyxy": [frame, 0, frame + 3, 8],
                        "centroid_uv": [frame + 1.5, 4],
                    }
                ],
                frame_index=frame,
                video_time_s=frame / 2,
            )
        state = build_given_that("current risk", memory, modality="rgb", image_size=(20, 20))
        visible = state["current_scene"]["visible_objects"][0]
        known = state["scene_memory"]["known_objects"][0]
        self.assertEqual([item["frame_index"] for item in visible["recent_poses"]], [2, 4])
        self.assertNotIn("recent_poses", known)

    def test_budget_drops_low_priority_objects(self) -> None:
        memory = SceneMemory()
        observations = []
        for index in range(30):
            observations.append(
                {
                    "object_id": f"object_{index:03d}",
                    "label": "thing",
                    "confidence": 0.5,
                    "bbox_xyxy": [index, index, index + 5, index + 5],
                    "centroid_uv": [index, index],
                    "description": "x" * 4000,
                }
            )
        memory.observe(observations)
        state = build_given_that("count", memory, modality="rgb", image_size=(100, 100))
        self.assertLessEqual(len(state["scene_memory"]["known_objects"]), 16)
        self.assertGreaterEqual(len(state["scene_memory"]["known_objects"]), 4)
        self.assertLessEqual(state["prompt_budget"]["serialized_chars"], 24000)


class SessionTests(unittest.TestCase):
    def test_ask_hides_given_that(self) -> None:
        captured = {}

        def client(state, questions):
            captured["state"] = state
            captured["questions"] = questions

            class Answer:
                choice = "yellow"
                confidence = 0.88
                probabilities = {"yellow": 0.88, "red": 0.12}

            class Response:
                model = "jev-latest"
                usage = None
                answers = {"answer": Answer()}

                def model_dump(self):
                    return {"model": self.model, "answers": {"answer": {"choice": "yellow"}}}

            return Response()

        sees = Sees(client=client)
        sees.memory.observe(
            [
                {
                    "object_id": "object_001",
                    "label": "bus",
                    "confidence": 0.9,
                    "bbox_xyxy": [1, 2, 3, 4],
                    "centroid_uv": [2, 3],
                    "attributes": {"color": "yellow"},
                }
            ]
        )
        sees.image_size = (640, 480)
        question = Choice(
            instructions="What color is the bus?",
            criteria={"yellow": None, "red": None},
        )
        result = sees.ask("What color is the bus?", {"answer": question})
        self.assertEqual(result.choice, "yellow")
        self.assertAlmostEqual(result.confidence, 0.88)
        self.assertNotIn("given_that", result.model_dump())
        self.assertIn("user_goal", captured["state"])
        self.assertIs(captured["questions"]["answer"], question)

    def test_state_is_ready_for_official_system_one_call(self) -> None:
        sees = Sees()
        sees.memory.observe(
            [
                {
                    "object_id": "object_001",
                    "label": "bus",
                    "confidence": 0.9,
                    "bbox_xyxy": [1, 2, 3, 4],
                    "centroid_uv": [2, 3],
                    "attributes": {"color": "yellow"},
                }
            ]
        )
        sees.image_size = (640, 480)

        state = sees.state("What color is the bus?")

        self.assertEqual(state["user_goal"], "What color is the bus?")
        self.assertEqual(
            state["current_scene"]["visible_objects"][0]["object_id"],
            "object_001",
        )
        self.assertNotIn("questions", state)
        self.assertNotIn("answers", state)


class ReportTests(unittest.TestCase):
    def test_risk_row_keeps_the_peak(self) -> None:
        log = JudgmentLog()
        person = {"object_id": "object_001", "label": "person", "bbox_xyxy": [1, 2, 3, 4]}

        class Answer:
            def __init__(self, noul: float):
                self.noul = noul
                self.choice = None

        log.add(10, [person], {"object_001": Answer(0.2)})
        log.add(16, [person], {"object_001": Answer(0.8)})
        row = log.rows()[0]
        self.assertEqual(row["object_id"], "object_001")
        self.assertEqual(row["risk"], 0.8)
        self.assertEqual(row["peak_frame"], 16)
        self.assertEqual([item["frame"] for item in log.export_timeline()], [10, 16])

    def test_video_result_is_grouped_by_frame(self) -> None:
        log = JudgmentLog()
        person = {"object_id": "object_001", "label": "person", "bbox_xyxy": [1, 2, 3, 4]}

        class Answer:
            noul = 0.65
            choice = None
            confidence = None

        log.add(12, [person], {"object_001": Answer()}, video_time_s=3.0)
        frame = log.export_frames()[0]
        self.assertEqual(frame["frame_index"], 12)
        self.assertEqual(frame["video_time_s"], 3.0)
        self.assertEqual(frame["objects"][0]["probability"], 0.65)

    def test_json_and_excel(self) -> None:
        import tempfile
        from pathlib import Path

        from openpyxl import load_workbook

        rows = [{"object_id": "object_001", "label": "person", "risk": 0.8, "peak_frame": 16}]
        timeline = [{"frame": 16, "object_id": "object_001", "risk": 0.8, "bbox_xyxy": [1, 2, 3, 4]}]
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            frames = [
                {
                    "frame_index": 16,
                    "video_time_s": 4.0,
                    "objects": [
                        {
                            "object_id": "object_001",
                            "label": "person",
                            "probability": 0.8,
                            "bbox_xyxy": [1, 2, 3, 4],
                        }
                    ],
                }
            ]
            write_json(root / "risk.json", "每个行人", rows, timeline, frames)
            write_excel(root / "risk.xlsx", rows, timeline, frames)
            payload = json.loads((root / "risk.json").read_text(encoding="utf-8"))
            self.assertEqual(payload["frames"][0]["frame_index"], 16)
            self.assertEqual(payload["rows"][0]["risk"], 0.8)
            self.assertEqual(payload["timeline"][0]["bbox_xyxy"], [1, 2, 3, 4])
            book = load_workbook(root / "risk.xlsx")
            self.assertEqual(book.sheetnames, ["Frames", "Summary", "Timeline"])
            self.assertEqual(book["Frames"]["A2"].value, 16)
            self.assertEqual(book["Timeline"]["A1"].value, "frame")


class RenderTests(unittest.TestCase):
    def test_dashboard_keeps_answer_rows_after_object_disappears(self) -> None:
        renderer = DashboardRenderer()

        class Answer:
            noul = 0.72
            confidence = None

        frame = np.zeros((80, 120, 3), dtype=np.uint8)
        tracks = [{"object_id": "object_001", "label": "person", "bbox_xyxy": [5, 5, 30, 70]}]
        first = renderer.render(
            frame,
            tracks,
            {"object_001": Answer()},
            frame_index=1,
            video_time_s=0.1,
            state_chars=1200,
        )
        second = renderer.render(
            frame,
            [],
            {},
            frame_index=2,
            video_time_s=0.2,
        )
        self.assertEqual(first.shape, second.shape)
        self.assertIn("object_001", renderer.entries)
        self.assertFalse(renderer.entries["object_001"]["visible"])


class CallTests(unittest.TestCase):
    def test_image_question_prints_a_choice_without_a_class_list(self) -> None:
        captured = {}

        def client(state, questions):
            captured["questions"] = questions

            class Answer:
                choice = "blue"
                confidence = 1.0
                probabilities = {"blue": 1.0}
                noul = None

            class Response:
                model = "jev-latest"
                usage = None
                answers = {"answer": Answer()}

                def model_dump(self):
                    return {"model": self.model, "answers": {"answer": {"choice": "blue"}}}

            return Response()

        sees = Sees(client=client)

        def observe(image, depth=None, *, intrinsics=None):
            return [{"object_id": "object_009", "label": "bus", "bbox_xyxy": [1, 1, 8, 8]}]

        sees.observe = observe
        image = np.zeros((16, 16, 3), dtype=np.uint8)
        question = Choice(
            instructions="What color is the bus?",
            criteria={"blue": None, "uncertain": None},
        )
        result = sees(image, "What color is the bus?", {"answer": question})
        self.assertIsInstance(captured["questions"]["answer"], Choice)
        self.assertEqual(result.rows[0]["choice"], "blue")
        self.assertEqual(result.choice, "blue")
        self.assertIn("blue", str(result))

    def test_video_builds_person_questions_from_tracks(self) -> None:
        import tempfile
        from pathlib import Path

        import cv2

        captured: list[list[str]] = []
        risks = iter((0.2, 0.8))

        def client(state, questions):
            captured.append(list(questions))
            risk = next(risks)

            class Answer:
                def __init__(self) -> None:
                    self.noul = risk
                    self.choice = None
                    self.confidence = None

            class Response:
                model = "jev-latest"
                usage = None
                answers = {key: Answer() for key in questions}

                def model_dump(self):
                    return {"model": self.model}

            return Response()

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            clip = root / "street.avi"
            writer = cv2.VideoWriter(str(clip), cv2.VideoWriter_fourcc(*"MJPG"), 4, (16, 16))
            self.assertTrue(writer.isOpened())
            for _ in range(2):
                writer.write(np.zeros((16, 16, 3), dtype=np.uint8))
            writer.release()
            sees = Sees(client=client)

            def observe(image, depth=None, *, intrinsics=None, **kwargs):
                return [
                    {"object_id": "object_001", "label": "person", "bbox_xyxy": [0, 0, 4, 8]},
                    {"object_id": "object_002", "label": "car", "bbox_xyxy": [5, 5, 12, 12]},
                ]

            sees.observe = observe

            def questions(tracks):
                return {
                    item["object_id"]: Noul(instructions="accident risk")
                    for item in tracks
                    if item["label"] == "person"
                }

            result = sees(
                clip,
                "帮我计算视频中每个行人发生 car accident 的风险",
                questions,
                every=1,
                json=root / "risk.json",
                excel=root / "risk.xlsx",
                save=root / "boxed.gif",
            )
            self.assertEqual(captured, [["object_001"], ["object_001"]])
            self.assertEqual(result.rows[0]["object_id"], "object_001")
            self.assertEqual(result.rows[0]["risk"], 0.8)
            self.assertEqual(result.rows[0]["peak_frame"], 1)
            self.assertEqual([frame["frame_index"] for frame in result.frames], [0, 1])
            self.assertEqual(result.frames[1]["objects"][0]["probability"], 0.8)
            self.assertTrue((root / "risk.json").is_file())
            self.assertTrue((root / "risk.xlsx").is_file())
            self.assertTrue((root / "boxed.gif").is_file())

    def test_connection_error_is_raised(self) -> None:
        import tempfile
        from pathlib import Path

        import cv2

        def client(state, questions):
            raise ConnectionError("[Errno 11002] getaddrinfo failed")

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            clip = root / "street.avi"
            writer = cv2.VideoWriter(str(clip), cv2.VideoWriter_fourcc(*"MJPG"), 4, (16, 16))
            self.assertTrue(writer.isOpened())
            writer.write(np.zeros((16, 16, 3), dtype=np.uint8))
            writer.release()
            sees = Sees(client=client)
            sees.observe = lambda image, depth=None, **kwargs: [
                {"object_id": "object_001", "label": "person", "bbox_xyxy": [0, 0, 4, 8]}
            ]
            with self.assertRaises(ConnectionError):
                sees(
                    clip,
                    "帮我计算视频中每个行人发生 car accident 的风险",
                    {"object_001": Noul(instructions="accident risk")},
                    every=1,
                    save=root / "boxed.gif",
                )
            self.assertFalse((root / "boxed.gif").is_file())

if __name__ == "__main__":
    unittest.main()
